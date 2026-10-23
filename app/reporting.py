"""Editable DOCX and PDF reports from investigator-approved snapshots."""

from io import BytesIO
from html import escape
import re
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image,
    KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from PIL import Image as PILImage, ImageDraw
from app.analytics import financial, statistics


def safe(value):
    return re.sub(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f]",
        lambda m: f"\\x{ord(m.group()):02x}",
        str(value),
    )


def sections(case, report, artifacts, sources):
    stats = statistics(artifacts)
    ai = (
        report.snapshot.get("analyses", []) if getattr(report, "snapshot", None) else []
    )
    return [
        (
            "I. Case Information",
            f"{case.reference}: {case.title}\nCategory: {case.category}\nOpened: {case.opened}\nDevice: {case.device}",
        ),
        ("II. Objective of Examination", case.description),
        (
            "III. Evidence Received",
            "\n".join(f"Source {e.id}: {e.filename} ({e.status})" for e in sources),
        ),
        (
            "IV. Examination Methodology",
            "Local parsing with original-field preservation; deterministic search and indicator extraction. Original uploads retained unchanged. Statistics in this report describe selected artifacts only.",
        ),
        (
            "V. Tools and Software Used",
            "CFAA 0.1.0; Python; FastAPI; SQLAlchemy; local Ollama cicc-agent:latest (when AI analysis requested).",
        ),
        (
            "VI. Evidence Integrity and Hash Values",
            "SHA-256 values recorded at import appear in the evidence inventory table below. Hashes are not a chain-of-custody certification.",
        ),
        ("VII. Examination Findings", report.body),
        (
            "VIII. Communication Analysis",
            "\n".join(f"{k}: {v}" for k, v in stats["categories"].items())
            + "\nDirections: "
            + str(stats["directions"])
            + "\nEndpoint occurrences: "
            + str(stats["contacts"]),
        ),
        (
            "IX. Timeline Analysis",
            "\n".join(
                f"Artifact {a.id}: {a.timestamp_original or 'missing'}; UTC: {a.timestamp_utc or 'unknown'}; {a.time_warning}"
                for a in sorted(artifacts, key=lambda a: a.timestamp_utc or "z")
            ),
        ),
        (
            "X. Relationship Analysis",
            "\n".join(
                f"Artifact {a.id}: {a.sender} -> {a.recipient}"
                for a in artifacts
                if a.sender and a.recipient
            )
            + "\nRelationships record endpoint associations only; they do not prove ownership or identity.",
        ),
        (
            "XI. Financial Indicators",
            "\n".join(
                f"Artifact {f['artifact_id']}: {f['kind']}: {f['value']}"
                for f in financial(artifacts)
            )
            + "\nPayment references do not establish completed transactions.",
        ),
        (
            "XII. AI-Assisted Analytical Observations",
            "AI observations below are unverified interpretations from the latest ten analyses available when this draft was saved. Source citation checks do not verify explanations.\n"
            + "\n\n".join(
                f"Analysis {a['id']}: {a['question']}\n"
                + "\n".join(
                    f"Artifact {f['artifact_id']}, source {f['source_id']}, record {f['source_ref']}: {f['excerpt']}\n{f['classification']}: {f['explanation']}\n{f['verification']}"
                    for f in a["result"].get("findings", [])
                )
                for a in ai
            ),
        ),
        (
            "XIII. Limitations",
            "Generic flat export adapters; no direct UFDR support. Missing timestamps and fields remain unknown. No certification or legal compliance assessment is implied. This report reflects its draft snapshot; review original evidence before relying on conclusions.",
        ),
        (
            "XIV. Conclusion",
            "The examination findings above reflect the investigator-approved narrative. Automated analysis makes no determination of guilt, identity, or criminal responsibility.",
        ),
        (
            "XV. Supporting Artifact References",
            "Selected evidence references are listed in the artifact table below.",
        ),
    ]


def chart(artifacts):
    counts = Counter(a.category for a in artifacts)
    canvas = PILImage.new("RGB", (1000, 60 + 30 * min(8, max(1, len(counts)))), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text((20, 10), "Selected artifacts by category", fill="#182d41")
    for i, (label, value) in enumerate(counts.most_common(8)):
        y = 40 + i * 30
        draw.text(
            (20, y),
            label.encode("ascii", errors="replace").decode()[:22],
            fill="#182d41",
        )
        draw.rectangle(
            (180, y, 180 + int(700 * value / max(counts.values())), y + 17),
            fill="#147d83",
        )
        draw.text((900, y), str(value), fill="#182d41")
    stream = BytesIO()
    canvas.save(stream, format="PNG")
    stream.seek(0)
    return stream


def generate(case, report, artifacts, sources, format):
    output = BytesIO()
    parts = sections(case, report, artifacts, sources)
    inventory = [["Source", "Filename", "SHA-256"]] + [
        [str(e.id), e.filename, e.digest] for e in sources
    ]
    refs = [
        ["Artifact / source / record", "Category / timestamp", "Evidence excerpt"]
    ] + [
        [
            f"{a.id} / {a.evidence_id} / {a.source_ref}",
            f"{a.category}\n{a.timestamp_original}\n{a.time_warning}",
            a.text,
        ]
        for a in artifacts
    ]
    if format == "docx":
        doc = Document()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(10)
        doc.add_heading("CICC Forensic Examination Report", 0)

        def table(data):
            t = doc.add_table(rows=0, cols=3)
            t.style = "Light Shading Accent 1"
            for row in data:
                for cell, value in zip(t.add_row().cells, row):
                    cell.text = safe(value)

        for title, body in parts:
            doc.add_heading(title, 1)
            for line in safe(body).splitlines():
                doc.add_paragraph(line)
            if title.startswith("VI."):
                table(inventory)
            if title.startswith("VIII.") and artifacts:
                doc.add_picture(chart(artifacts), width=Inches(6))
            if title.startswith("XV."):
                table(refs)
        doc.save(output)
    else:
        styles = getSampleStyleSheet()
        styles["BodyText"].wordWrap = "CJK"
        styles["BodyText"].fontSize = 9.5
        styles["BodyText"].leading = 13
        styles["BodyText"].spaceBefore = 0
        styles["BodyText"].spaceAfter = 3
        styles["Heading2"].fontSize = 12
        styles["Heading2"].leading = 15
        styles["Heading2"].spaceBefore = 10
        styles["Heading2"].spaceAfter = 5
        story = [Paragraph("CICC Forensic Examination Report", styles["Title"])]

        def para(value):
            return Paragraph(
                escape(safe(value)).replace("\n", "<br/>"), styles["BodyText"]
            )

        def table(data):
            t = Table(
                [[para(c) for c in row] for row in data],
                colWidths=[100, 140, 270],
                repeatRows=1,
                splitByRow=1,
                splitInRow=1,
            )
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eff4")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd6e0")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            return t

        for title, body in parts:
            block_start = len(story)
            story.append(Paragraph(escape(title), styles["Heading2"]))
            for line in safe(body).splitlines():
                story.append(para(line))
            if title.startswith("VI."):
                story.append(table(inventory))
            if title.startswith("VIII.") and artifacts:
                story.append(
                    Image(
                        chart(artifacts),
                        width=480,
                        height=480
                        * (
                            60
                            + 30
                            * min(8, max(1, len(set(a.category for a in artifacts))))
                        )
                        / 1000,
                    )
                )
            if title.startswith("XV."):
                story.append(table(refs))
            story.append(Spacer(1, 8))
            if title.startswith(("VIII.", "XV.")):
                story[block_start:] = [KeepTogether(story[block_start:])]
        SimpleDocTemplate(output, leftMargin=42, rightMargin=42).build(story)
    return output.getvalue()
