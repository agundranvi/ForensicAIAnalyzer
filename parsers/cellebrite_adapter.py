import csv, json, zipfile
from datetime import datetime, timezone
from pathlib import Path
from defusedxml import ElementTree
from openpyxl import load_workbook

LIMIT = 50000


def records(path: Path, extension: str):
    if extension == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            if reader.fieldnames and len(set(reader.fieldnames)) != len(
                reader.fieldnames
            ):
                raise ValueError("Duplicate CSV headers")
            for i, row in enumerate(reader, 2):
                if None in row:
                    raise ValueError(f"Record {i}: more fields than headers")
                yield str(i), row
    elif extension == ".json":
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(value, dict):
            value = value.get("artifacts", value.get("records", [value]))
        if not isinstance(value, list):
            raise ValueError("JSON must contain a record list")
        for i, row in enumerate(value, 1):
            if not isinstance(row, dict):
                raise ValueError("Every JSON record must be an object")
            yield str(i), row
    elif extension == ".xml":
        root = ElementTree.parse(path).getroot()
        for i, element in enumerate(list(root), 1):
            yield (
                str(i),
                {
                    "_tag": element.tag,
                    **element.attrib,
                    **{child.tag: "".join(child.itertext()) for child in element},
                    "_xml": ElementTree.tostring(element, encoding="unicode"),
                },
            )
    elif extension == ".xlsx":
        with zipfile.ZipFile(path) as z:
            if (
                len(z.infolist()) > 10000
                or sum(x.file_size for x in z.infolist()) > 100 * 1024 * 1024
            ):
                raise ValueError("Workbook expansion limit exceeded")
            if any(
                ".." in Path(x.filename).parts or x.filename.startswith("/")
                for x in z.infolist()
            ):
                raise ValueError("Unsafe archive path")
        source = path.open("rb")
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book:
                rows = sheet.iter_rows(values_only=True)
                headers = [
                    str(v) if v is not None else f"column_{i}"
                    for i, v in enumerate(next(rows, []))
                ]
                if len(set(headers)) != len(headers):
                    raise ValueError("Duplicate spreadsheet headers")
                for i, row in enumerate(rows, 2):
                    yield f"{sheet.title}:{i}", dict(zip(headers, row))
        finally:
            book.close()
            source.close()
    else:
        raise ValueError(
            "Supported formats: CSV, XLSX, JSON, XML. UFDR and HTML are not supported."
        )


def normalize(row):
    raw = json.loads(json.dumps(row, default=str))
    fields = {
        str(k).lower().replace(" ", "_"): str(v) if v is not None else ""
        for k, v in raw.items()
    }

    def pick(*keys):
        return next((fields[k] for k in keys if fields.get(k)), "")

    original = pick("timestamp", "date", "time", "datetime", "date/time")
    utc, warning = "", "Missing timestamp"
    if original:
        try:
            dt = datetime.fromisoformat(original.replace("Z", "+00:00"))
            if dt.tzinfo:
                utc, warning = dt.astimezone(timezone.utc).isoformat(), ""
            else:
                warning = "Timezone unknown; UTC conversion withheld"
        except ValueError:
            warning = "Unrecognized timestamp; original preserved"
    category = pick("category", "type", "artifact_type", "_tag") or "unsupported"
    return dict(
        category=category.lower()[:100],
        sender=pick("sender", "from", "source", "phone", "name"),
        recipient=pick("recipient", "to", "destination"),
        direction=pick("direction")[:100],
        application=pick("application", "app")[:100],
        timestamp_original=original[:255],
        timestamp_utc=utc,
        time_warning=warning,
        text=pick("text", "body", "message", "content", "url", "description")
        or json.dumps(raw, ensure_ascii=False),
        raw=raw,
    )
